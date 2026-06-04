#!/usr/bin/env python3
"""
交易复盘日志系统
功能：图片识别交割单 / 手动录入 / 复盘笔记 / 导出Excel
运行方式：python trade_journal.py
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import json, os, base64, datetime, re
from pathlib import Path
import openpyxl
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

# ── Config ──────────────────────────────────────────────────────────────────
DATA_FILE   = Path.home() / ".trade_journal_data.json"
API_KEY_FILE = Path.home() / ".trade_journal_api_key.txt"

COLORS = {
    "bg":       "#0e0e14",
    "panel":    "#15151e",
    "border":   "#2a2a3a",
    "accent":   "#e84040",
    "green":    "#2ecc71",
    "red":      "#e74c3c",
    "text":     "#e0e0e0",
    "muted":    "#888899",
    "input_bg": "#1e1e2a",
    "btn":      "#252535",
    "btn_hover":"#303045",
}

DIRECTIONS  = ["买入 (Buy)", "卖出 (Sell)"]
RESULTS     = ["盈利", "亏损", "保本"]
EMOTIONS    = ["冷静", "贪婪", "恐惧", "犹豫", "冲动", "FOMO"]
GRADE_OPTS  = ["A — 完美执行", "B — 基本合理", "C — 有瑕疵", "D — 错误判断", "F — 严重失误"]

# ── Data Layer ───────────────────────────────────────────────────────────────

def load_data():
    if DATA_FILE.exists():
        try:
            return json.loads(DATA_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {"trades": [], "settings": {}}

def save_data(data):
    DATA_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2),
                         encoding="utf-8")

def new_trade():
    return {
        "id": datetime.datetime.now().strftime("%Y%m%d%H%M%S%f"),
        "date": datetime.date.today().isoformat(),
        "time": datetime.datetime.now().strftime("%H:%M"),
        "symbol": "",
        "direction": "买入 (Buy)",
        "quantity": "",
        "price": "",
        "amount": "",
        "stop_loss": "",
        "take_profit": "",
        "result": "盈利",
        "pnl": "",
        "commission": "",
        "emotion": "冷静",
        "grade": "B — 基本合理",
        "reason_entry": "",
        "reason_exit": "",
        "review": "",
        "lesson": "",
        "image_path": "",
        "tags": "",
    }

# ── Claude API Image Recognition ─────────────────────────────────────────────

def ocr_trade_image(image_path: str, api_key: str) -> dict:
    """Send image to Claude API, extract trade details."""
    import urllib.request, urllib.error
    
    path = Path(image_path)
    suffix = path.suffix.lower()
    mime = {".jpg":"image/jpeg",".jpeg":"image/jpeg",
            ".png":"image/png",".gif":"image/gif",".webp":"image/webp"
            }.get(suffix, "image/jpeg")
    
    b64 = base64.b64encode(path.read_bytes()).decode()
    
    prompt = """请仔细分析这张交易截图/交割单，提取以下信息，严格用JSON格式返回，不要加任何说明文字：
{
  "date": "YYYY-MM-DD格式日期，如果有的话",
  "time": "HH:MM格式时间，如果有的话",
  "symbol": "交易品种代码，如AAPL、NOK、BTC等",
  "direction": "买入或卖出",
  "quantity": "数量，纯数字",
  "price": "成交价格，纯数字",
  "amount": "成交金额，纯数字",
  "commission": "手续费，纯数字，如果有的话",
  "pnl": "盈亏金额，如果有的话，负数表示亏损",
  "notes": "其他备注信息"
}
如果某字段无法识别，填null。只返回JSON。"""

    payload = json.dumps({
        "model": "claude-sonnet-4-20250514",
        "max_tokens": 500,
        "messages": [{
            "role": "user",
            "content": [
                {"type": "image",
                 "source": {"type": "base64", "media_type": mime, "data": b64}},
                {"type": "text", "text": prompt}
            ]
        }]
    }).encode()

    req = urllib.request.Request(
        "https://api.anthropic.com/v1/messages",
        data=payload,
        headers={
            "Content-Type": "application/json",
            "x-api-key": api_key,
            "anthropic-version": "2023-06-01",
        },
        method="POST"
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        body = json.loads(resp.read())
    
    text = body["content"][0]["text"].strip()
    text = re.sub(r"^```[a-z]*\n?", "", text)
    text = re.sub(r"\n?```$", "", text)
    return json.loads(text)

# ── Excel Export ──────────────────────────────────────────────────────────────

def export_to_excel(trades: list, filepath: str):
    wb = openpyxl.Workbook()

    # ── Sheet 1: 交易记录 ──
    ws1 = wb.active
    ws1.title = "交易记录"

    header_fill  = PatternFill("solid", fgColor="1a1a2e")
    header_font  = Font(bold=True, color="e0e0e0", size=10)
    buy_fill     = PatternFill("solid", fgColor="0d2b1a")
    sell_fill    = PatternFill("solid", fgColor="2b0d0d")
    profit_font  = Font(color="2ecc71", bold=True)
    loss_font    = Font(color="e74c3c", bold=True)
    border_thin  = Border(
        left=Side(style="thin", color="2a2a3a"),
        right=Side(style="thin", color="2a2a3a"),
        top=Side(style="thin", color="2a2a3a"),
        bottom=Side(style="thin", color="2a2a3a"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left   = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    headers = [
        "日期","时间","品种","方向","数量","价格","金额($)",
        "止损价","止盈价","手续费($)","盈亏($)","结果","情绪状态","评级",
        "入场理由","离场理由","复盘总结","经验教训","标签"
    ]
    col_widths = [12,8,10,10,8,10,12,10,10,10,12,8,10,16,30,30,40,40,20]

    ws1.row_dimensions[1].height = 30
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws1.cell(1, ci, h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = border_thin
        ws1.column_dimensions[get_column_letter(ci)].width = w

    ws1.freeze_panes = "A2"

    grade_colors = {
        "A": "1a4a1a", "B": "1a3a2a", "C": "3a3a0a",
        "D": "3a1a0a", "F": "4a0a0a"
    }

    for ri, t in enumerate(trades, 2):
        ws1.row_dimensions[ri].height = 20
        is_buy  = "买" in t.get("direction", "")
        row_fill = buy_fill if is_buy else sell_fill
        grade_letter = t.get("grade", "B")[0]
        g_color = grade_colors.get(grade_letter, "1e1e2a")

        def c(col_idx, val, fmt=None):
            cell = ws1.cell(ri, col_idx, val)
            cell.fill   = row_fill
            cell.border = border_thin
            cell.alignment = center
            if fmt:
                cell.number_format = fmt
            return cell

        c(1,  t.get("date",""))
        c(2,  t.get("time",""))
        c(3,  t.get("symbol","").upper())
        dir_cell = c(4, "▲ 买入" if is_buy else "▼ 卖出")
        dir_cell.font = Font(color="4fc3f7" if is_buy else "ef5350", bold=True)

        for ci2, key in zip([5,6,7,8,9,10], ["quantity","price","amount","stop_loss","take_profit","commission"]):
            try:    c(ci2, float(t.get(key,"") or 0), "#,##0.00")
            except: c(ci2, t.get(key,""))

        pnl_raw = t.get("pnl","")
        try:
            pnl_val = float(pnl_raw or 0)
            pnl_cell = c(11, pnl_val, '+#,##0.00;-#,##0.00;"-"')
            pnl_cell.font = profit_font if pnl_val > 0 else (loss_font if pnl_val < 0 else Font())
        except:
            c(11, pnl_raw)

        result = t.get("result","")
        res_cell = c(12, result)
        res_cell.font = Font(color="2ecc71" if result=="盈利" else ("e74c3c" if result=="亏损" else "aaaaaa"), bold=True)

        c(13, t.get("emotion",""))
        grade_cell = c(14, t.get("grade",""))
        grade_cell.fill = PatternFill("solid", fgColor=g_color)
        grade_cell.font = Font(color="e0e0e0", bold=True, size=9)

        for ci2, key in zip([15,16,17,18,19],
                            ["reason_entry","reason_exit","review","lesson","tags"]):
            cell = ws1.cell(ri, ci2, t.get(key,""))
            cell.fill   = row_fill
            cell.border = border_thin
            cell.alignment = left

    # ── Sheet 2: 统计分析 ──
    ws2 = wb.create_sheet("统计分析")
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 18
    ws2.column_dimensions["D"].width = 28

    stat_header = Font(bold=True, color="ffffff", size=11)
    stat_fill   = PatternFill("solid", fgColor="1a1a2e")
    val_fill    = PatternFill("solid", fgColor="111120")

    def stat_row(r, label, formula_or_val, note=""):
        a = ws2.cell(r, 1, label)
        a.font = Font(bold=True, color="aaaacc", size=10)
        a.fill = val_fill
        a.border = border_thin
        a.alignment = left

        b = ws2.cell(r, 2, formula_or_val)
        b.fill = val_fill
        b.border = border_thin
        b.alignment = center

        if note:
            d = ws2.cell(r, 4, note)
            d.font = Font(color="666688", size=9, italic=True)
            d.alignment = left

        ws2.row_dimensions[r].height = 22

    # Title
    ws2["A1"] = "📊 交易统计分析"
    ws2["A1"].font = Font(bold=True, color="e0e0e0", size=14)
    ws2["A1"].fill = PatternFill("solid", fgColor="0d0d1a")
    ws2.merge_cells("A1:D1")
    ws2.row_dimensions[1].height = 32

    data_sheet = "交易记录"
    last_row = len(trades) + 1
    n = last_row

    stat_row(2, "总交易笔数",  f"=COUNTA('{data_sheet}'!A2:A{n})")
    stat_row(3, "总盈亏 ($)",  f"=IFERROR(SUM('{data_sheet}'!K2:K{n}),0)", "所有已实现PNL之和")
    stat_row(4, "盈利笔数",    f"=COUNTIF('{data_sheet}'!L2:L{n},\"盈利\")")
    stat_row(5, "亏损笔数",    f"=COUNTIF('{data_sheet}'!L2:L{n},\"亏损\")")
    stat_row(6, "胜率",        f"=IFERROR(B4/B2,0)", "盈利笔数/总笔数")
    ws2["B6"].number_format = "0.0%"

    stat_row(7,  "平均盈利 ($)", f"=IFERROR(AVERAGEIF('{data_sheet}'!L2:L{n},\"盈利\",'{data_sheet}'!K2:K{n}),0)")
    stat_row(8,  "平均亏损 ($)", f"=IFERROR(AVERAGEIF('{data_sheet}'!L2:L{n},\"亏损\",'{data_sheet}'!K2:K{n}),0)")
    stat_row(9,  "盈亏比",       "=IFERROR(ABS(B7/B8),0)", "平均盈利/平均亏损绝对值")
    stat_row(10, "最大单笔盈利", f"=IFERROR(MAX('{data_sheet}'!K2:K{n}),0)")
    stat_row(11, "最大单笔亏损", f"=IFERROR(MIN('{data_sheet}'!K2:K{n}),0)")
    stat_row(12, "总手续费 ($)", f"=IFERROR(SUM('{data_sheet}'!J2:J{n}),0)")

    ws2["A14"] = "情绪分布"
    ws2["A14"].font = Font(bold=True, color="aaaacc", size=10)
    ws2["A14"].fill = PatternFill("solid", fgColor="1a1a2e")
    for r2, emo in enumerate(EMOTIONS, 15):
        ws2.cell(r2, 1, emo).fill = val_fill
        f = f"=COUNTIF('{data_sheet}'!M2:M{n},\"{emo}\")"
        ws2.cell(r2, 2, f).fill = val_fill
        ws2.cell(r2, 2).alignment = center

    ws2["C14"] = "评级分布"
    ws2["C14"].font = Font(bold=True, color="aaaacc", size=10)
    ws2["C14"].fill = PatternFill("solid", fgColor="1a1a2e")
    for r2, g in enumerate(["A","B","C","D","F"], 15):
        ws2.cell(r2, 3, f"Grade {g}").fill = val_fill
        ws2.cell(r2, 4, f"=COUNTIF('{data_sheet}'!N2:N{n},\"{g}*\")").fill = val_fill
        ws2.cell(r2, 4).alignment = center
        ws2.row_dimensions[r2].height = 20

    # Format stat values
    for r2 in range(2, 13):
        cell = ws2.cell(r2, 2)
        if r2 in [3, 7, 8, 10, 11, 12]:
            cell.number_format = '[Blue]+#,##0.00;[Red]-#,##0.00;"-"'
        elif r2 == 9:
            cell.number_format = "0.00x"

    # ── Sheet 3: 每日PNL ──
    ws3 = wb.create_sheet("每日盈亏")
    ws3["A1"] = "日期"
    ws3["B1"] = "当日盈亏 ($)"
    ws3["C1"] = "累计盈亏 ($)"
    ws3["D1"] = "交易笔数"
    for cell in ws3["A1:D1"][0]:
        cell.font  = header_font
        cell.fill  = header_fill
        cell.alignment = center
        cell.border = border_thin
    ws3.column_dimensions["A"].width = 14
    ws3.column_dimensions["B"].width = 16
    ws3.column_dimensions["C"].width = 16
    ws3.column_dimensions["D"].width = 12

    # Group trades by date
    from collections import defaultdict
    daily = defaultdict(lambda: {"pnl": 0.0, "count": 0})
    for t in trades:
        d = t.get("date","")
        if not d: continue
        try: daily[d]["pnl"] += float(t.get("pnl","") or 0)
        except: pass
        daily[d]["count"] += 1

    cum = 0.0
    for ri3, (day, info) in enumerate(sorted(daily.items()), 2):
        cum += info["pnl"]
        ws3.cell(ri3, 1, day).alignment = center
        pnl_c = ws3.cell(ri3, 2, round(info["pnl"],2))
        pnl_c.number_format = '+#,##0.00;-#,##0.00;"-"'
        pnl_c.font = Font(color="2ecc71" if info["pnl"]>=0 else "e74c3c")
        cum_c = ws3.cell(ri3, 3, round(cum,2))
        cum_c.number_format = '+#,##0.00;-#,##0.00;"-"'
        cum_c.font = Font(color="2ecc71" if cum>=0 else "e74c3c", bold=True)
        ws3.cell(ri3, 4, info["count"]).alignment = center
        for ci2 in range(1,5):
            ws3.cell(ri3, ci2).border = border_thin

    # Add chart if enough data
    if len(daily) >= 2:
        chart = LineChart()
        chart.title = "累计盈亏曲线"
        chart.style = 10
        chart.y_axis.title = "盈亏 ($)"
        chart.x_axis.title = "日期"
        chart.height = 12
        chart.width  = 22
        data_ref = Reference(ws3, min_col=3, min_row=1, max_row=len(daily)+1)
        chart.add_data(data_ref, titles_from_data=True)
        ws3.add_chart(chart, "F2")

    wb.save(filepath)
    return filepath

# ── Main GUI ──────────────────────────────────────────────────────────────────

class TradeJournalApp:
    def __init__(self, root):
        self.root = root
        self.root.title("交易复盘日志 | Trade Journal")
        self.root.geometry("1200x780")
        self.root.minsize(900, 600)
        self.root.configure(bg=COLORS["bg"])

        self.data    = load_data()
        self.trades  = self.data["trades"]
        self.current = new_trade()
        self.edit_idx = None  # None = new trade

        self._build_ui()
        self._refresh_list()

    # ── UI construction ──────────────────────────────────────────────────────

    def _build_ui(self):
        # Top bar
        topbar = tk.Frame(self.root, bg="#0a0a10", height=48)
        topbar.pack(fill="x", side="top")
        topbar.pack_propagate(False)

        tk.Label(topbar, text="📈  交易复盘日志", bg="#0a0a10",
                 fg="#e84040", font=("Georgia", 15, "bold")).pack(side="left", padx=20, pady=10)

        tk.Button(topbar, text="⚙ API设置",  command=self._settings_dialog,
                  **self._btn_style()).pack(side="right", padx=6, pady=8)
        tk.Button(topbar, text="📤 导出 Excel", command=self._export_excel,
                  bg="#1a4a1a", fg="#2ecc71", font=("Helvetica",10,"bold"),
                  relief="flat", cursor="hand2", padx=14).pack(side="right", padx=2, pady=8)

        # Main layout: left list + right form
        main = tk.Frame(self.root, bg=COLORS["bg"])
        main.pack(fill="both", expand=True, padx=12, pady=8)

        # ── LEFT: trade list ──
        left = tk.Frame(main, bg=COLORS["panel"], width=310)
        left.pack(side="left", fill="y", padx=(0,8))
        left.pack_propagate(False)

        tk.Label(left, text="交易记录", bg=COLORS["panel"],
                 fg=COLORS["muted"], font=("Helvetica",10,"bold")).pack(anchor="w", padx=12, pady=(10,4))

        # Search bar
        search_fr = tk.Frame(left, bg=COLORS["panel"])
        search_fr.pack(fill="x", padx=10, pady=(0,6))
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *_: self._refresh_list())
        tk.Entry(search_fr, textvariable=self.search_var, bg=COLORS["input_bg"],
                 fg=COLORS["text"], insertbackground=COLORS["text"],
                 relief="flat", font=("Helvetica",10), highlightthickness=1,
                 highlightbackground=COLORS["border"]).pack(fill="x", ipady=5)

        list_fr = tk.Frame(left, bg=COLORS["panel"])
        list_fr.pack(fill="both", expand=True, padx=10, pady=(0,8))

        sb = tk.Scrollbar(list_fr, bg=COLORS["panel"], troughcolor=COLORS["border"])
        sb.pack(side="right", fill="y")
        self.listbox = tk.Listbox(list_fr, bg=COLORS["input_bg"], fg=COLORS["text"],
                                   selectbackground=COLORS["accent"],
                                   font=("Courier",10), relief="flat",
                                   yscrollcommand=sb.set, activestyle="none",
                                   highlightthickness=0)
        self.listbox.pack(fill="both", expand=True)
        sb.config(command=self.listbox.yview)
        self.listbox.bind("<<ListboxSelect>>", self._on_select)

        btn_row = tk.Frame(left, bg=COLORS["panel"])
        btn_row.pack(fill="x", padx=10, pady=(0,10))
        tk.Button(btn_row, text="＋ 新建", command=self._new_trade,
                  **self._btn_style(accent=True)).pack(side="left", fill="x", expand=True, padx=(0,4))
        tk.Button(btn_row, text="🗑 删除", command=self._delete_trade,
                  **self._btn_style()).pack(side="left", fill="x", expand=True)

        # ── RIGHT: form ──
        right = tk.Frame(main, bg=COLORS["bg"])
        right.pack(side="left", fill="both", expand=True)

        # Notebook tabs
        style = ttk.Style()
        style.theme_use("default")
        style.configure("Dark.TNotebook",
                        background=COLORS["bg"], borderwidth=0)
        style.configure("Dark.TNotebook.Tab",
                        background=COLORS["btn"], foreground=COLORS["muted"],
                        padding=[14,6], font=("Helvetica",10,"bold"))
        style.map("Dark.TNotebook.Tab",
                  background=[("selected", COLORS["panel"])],
                  foreground=[("selected", COLORS["text"])])

        self.nb = ttk.Notebook(right, style="Dark.TNotebook")
        self.nb.pack(fill="both", expand=True)

        self._build_tab_basic()
        self._build_tab_review()
        self._build_tab_image()

        # Save button
        save_row = tk.Frame(right, bg=COLORS["bg"])
        save_row.pack(fill="x", pady=8)
        tk.Button(save_row, text="💾  保存这笔交易", command=self._save_trade,
                  bg=COLORS["accent"], fg="white", font=("Helvetica",12,"bold"),
                  relief="flat", cursor="hand2", padx=20, pady=8).pack(padx=4)

    def _build_tab_basic(self):
        fr = tk.Frame(self.nb, bg=COLORS["panel"])
        self.nb.add(fr, text="  基本信息  ")

        canvas = tk.Canvas(fr, bg=COLORS["panel"], highlightthickness=0)
        vsb = tk.Scrollbar(fr, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        inner = tk.Frame(canvas, bg=COLORS["panel"])
        canvas_win = canvas.create_window((0,0), window=inner, anchor="nw")
        inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(canvas_win, width=e.width))

        self.basic_vars = {}
        fields = [
            ("日期",   "date",      "entry"),
            ("时间",   "time",      "entry"),
            ("品种",   "symbol",    "entry"),
            ("方向",   "direction", "combo", DIRECTIONS),
            ("数量",   "quantity",  "entry"),
            ("成交价", "price",     "entry"),
            ("金额($)","amount",    "entry"),
            ("止损价", "stop_loss", "entry"),
            ("止盈价", "take_profit","entry"),
            ("手续费($)","commission","entry"),
            ("盈亏($)", "pnl",      "entry"),
            ("结果",   "result",    "combo", RESULTS),
            ("情绪",   "emotion",   "combo", EMOTIONS),
            ("评级",   "grade",     "combo", GRADE_OPTS),
            ("标签",   "tags",      "entry"),
        ]
        cols = 3
        for i, row_def in enumerate(fields):
            label, key = row_def[0], row_def[1]
            kind = row_def[2]
            r, c = divmod(i, cols)
            tk.Label(inner, text=label, bg=COLORS["panel"],
                     fg=COLORS["muted"], font=("Helvetica",9)).grid(
                row=r*2, column=c, sticky="w", padx=(14,4), pady=(8,0))

            if kind == "entry":
                var = tk.StringVar()
                w = tk.Entry(inner, textvariable=var, bg=COLORS["input_bg"],
                             fg=COLORS["text"], insertbackground=COLORS["text"],
                             relief="flat", font=("Helvetica",10),
                             highlightthickness=1, highlightbackground=COLORS["border"])
                w.grid(row=r*2+1, column=c, sticky="ew", padx=(14,4), pady=(0,2), ipady=5)
            else:
                opts = row_def[3]
                var = tk.StringVar(value=opts[0])
                w = ttk.Combobox(inner, textvariable=var, values=opts,
                                 state="readonly", font=("Helvetica",10))
                w.grid(row=r*2+1, column=c, sticky="ew", padx=(14,4), pady=(0,2), ipady=3)

            self.basic_vars[key] = var
            inner.columnconfigure(c, weight=1)

    def _build_tab_review(self):
        fr = tk.Frame(self.nb, bg=COLORS["panel"])
        self.nb.add(fr, text="  复盘笔记  ")

        self.review_vars = {}
        fields = [
            ("入场理由",   "reason_entry", 3),
            ("离场理由",   "reason_exit",  3),
            ("复盘总结",   "review",       5),
            ("经验教训",   "lesson",       4),
        ]
        for label, key, h in fields:
            tk.Label(fr, text=label, bg=COLORS["panel"],
                     fg=COLORS["muted"], font=("Helvetica",9,"bold")).pack(
                anchor="w", padx=14, pady=(10,2))
            txt = scrolledtext.ScrolledText(
                fr, height=h, bg=COLORS["input_bg"],
                fg=COLORS["text"], insertbackground=COLORS["text"],
                relief="flat", font=("Helvetica",10), wrap="word",
                highlightthickness=1, highlightbackground=COLORS["border"])
            txt.pack(fill="x", padx=14, pady=(0,2))
            self.review_vars[key] = txt

    def _build_tab_image(self):
        fr = tk.Frame(self.nb, bg=COLORS["panel"])
        self.nb.add(fr, text="  图片识别  ")

        tk.Label(fr, text="上传交割单截图，AI自动识别交易信息",
                 bg=COLORS["panel"], fg=COLORS["muted"],
                 font=("Helvetica",10)).pack(pady=(20,8))

        btn_row = tk.Frame(fr, bg=COLORS["panel"])
        btn_row.pack(pady=6)
        tk.Button(btn_row, text="📁 选择图片", command=self._choose_image,
                  **self._btn_style()).pack(side="left", padx=6)
        tk.Button(btn_row, text="🔍 AI识别", command=self._ocr_image,
                  **self._btn_style(accent=True)).pack(side="left", padx=6)

        self.img_path_var = tk.StringVar(value="未选择图片")
        tk.Label(fr, textvariable=self.img_path_var, bg=COLORS["panel"],
                 fg=COLORS["muted"], font=("Helvetica",9), wraplength=400).pack(pady=4)

        self.ocr_status = tk.Label(fr, text="", bg=COLORS["panel"],
                                    fg=COLORS["green"], font=("Helvetica",9,"bold"))
        self.ocr_status.pack(pady=4)

        tk.Label(fr, text="识别结果预览（可编辑后点保存）：",
                 bg=COLORS["panel"], fg=COLORS["muted"],
                 font=("Helvetica",9)).pack(anchor="w", padx=14, pady=(12,2))
        self.ocr_preview = scrolledtext.ScrolledText(
            fr, height=8, bg=COLORS["input_bg"],
            fg=COLORS["text"], insertbackground=COLORS["text"],
            relief="flat", font=("Courier",10),
            highlightthickness=1, highlightbackground=COLORS["border"])
        self.ocr_preview.pack(fill="x", padx=14)

        tk.Label(fr, text="⚠ 需要在设置中填入 Claude API Key 才能使用AI识别",
                 bg=COLORS["panel"], fg="#886633",
                 font=("Helvetica",8,"italic")).pack(pady=(8,4))

    # ── Event handlers ───────────────────────────────────────────────────────

    def _btn_style(self, accent=False):
        return dict(
            bg=COLORS["accent"] if accent else COLORS["btn"],
            fg="white",
            font=("Helvetica",10,"bold") if accent else ("Helvetica",10),
            relief="flat", cursor="hand2", padx=10, pady=4,
            activebackground=COLORS["red"] if accent else COLORS["btn_hover"],
            activeforeground="white",
        )

    def _refresh_list(self):
        self.listbox.delete(0, "end")
        q = self.search_var.get().lower() if hasattr(self,"search_var") else ""
        self._list_indices = []
        for i, t in enumerate(self.trades):
            txt = f"{t.get('date','')} {t.get('symbol',''):>6}  {'▲' if '买' in t.get('direction','') else '▼'}  {t.get('pnl','')}"
            if q and q not in txt.lower() and q not in t.get("review","").lower():
                continue
            self.listbox.insert("end", txt)
            self._list_indices.append(i)
            pnl = t.get("pnl","")
            try:
                color = COLORS["green"] if float(pnl) > 0 else COLORS["red"] if float(pnl) < 0 else COLORS["text"]
            except:
                color = COLORS["text"]
            self.listbox.itemconfig("end", fg=color)

    def _on_select(self, event=None):
        sel = self.listbox.curselection()
        if not sel: return
        self.edit_idx = self._list_indices[sel[0]]
        self.current  = dict(self.trades[self.edit_idx])
        self._load_to_form()

    def _load_to_form(self):
        t = self.current
        for key, var in self.basic_vars.items():
            var.set(t.get(key,""))
        for key, widget in self.review_vars.items():
            widget.delete("1.0","end")
            widget.insert("1.0", t.get(key,""))
        if t.get("image_path"):
            self.img_path_var.set(t["image_path"])

    def _new_trade(self):
        self.edit_idx = None
        self.current  = new_trade()
        self._load_to_form()
        self.nb.select(0)
        self.listbox.selection_clear(0,"end")

    def _save_trade(self):
        t = self.current
        for key, var in self.basic_vars.items():
            t[key] = var.get().strip()
        for key, widget in self.review_vars.items():
            t[key] = widget.get("1.0","end").strip()
        t["image_path"] = self.img_path_var.get() if self.img_path_var.get() != "未选择图片" else ""

        if not t.get("symbol"):
            messagebox.showwarning("提示", "请填写交易品种")
            return

        if self.edit_idx is None:
            self.trades.append(t)
        else:
            self.trades[self.edit_idx] = t

        self.data["trades"] = self.trades
        save_data(self.data)
        self._refresh_list()
        messagebox.showinfo("✓", f"已保存：{t['symbol']} {t['date']}")

    def _delete_trade(self):
        sel = self.listbox.curselection()
        if not sel:
            messagebox.showwarning("提示", "请先选择一笔交易")
            return
        idx = self._list_indices[sel[0]]
        t   = self.trades[idx]
        if messagebox.askyesno("确认删除", f"删除 {t.get('symbol','')} {t.get('date','')}？"):
            self.trades.pop(idx)
            save_data(self.data)
            self._refresh_list()
            self._new_trade()

    def _choose_image(self):
        path = filedialog.askopenfilename(
            title="选择交割单图片",
            filetypes=[("图片文件","*.png *.jpg *.jpeg *.webp *.gif"),("所有文件","*.*")]
        )
        if path:
            self.img_path_var.set(path)
            self.current["image_path"] = path

    def _ocr_image(self):
        path = self.img_path_var.get()
        if path == "未选择图片" or not Path(path).exists():
            messagebox.showwarning("提示", "请先选择图片文件")
            return
        api_key = self._get_api_key()
        if not api_key:
            messagebox.showwarning("需要API Key", "请在设置中填入 Claude API Key")
            self._settings_dialog()
            return
        self.ocr_status.config(text="正在识别中...", fg="#f39c12")
        self.root.update()
        try:
            result = ocr_trade_image(path, api_key)
            # Fill form fields from result
            mapping = {
                "date":"date","time":"time","symbol":"symbol",
                "direction":"direction","quantity":"quantity",
                "price":"price","amount":"amount",
                "commission":"commission","pnl":"pnl",
            }
            for api_key2, form_key in mapping.items():
                val = result.get(api_key2)
                if val is not None and str(val) not in ("null","None",""):
                    if form_key == "direction":
                        val_str = str(val)
                        if "买" in val_str or "buy" in val_str.lower():
                            self.basic_vars[form_key].set(DIRECTIONS[0])
                        else:
                            self.basic_vars[form_key].set(DIRECTIONS[1])
                    else:
                        self.basic_vars[form_key].set(str(val))

            pretty = json.dumps(result, ensure_ascii=False, indent=2)
            self.ocr_preview.delete("1.0","end")
            self.ocr_preview.insert("1.0", pretty)
            self.ocr_status.config(text="✓ 识别成功！请检查并补充信息后保存", fg=COLORS["green"])
        except Exception as e:
            self.ocr_status.config(text=f"✗ 识别失败：{e}", fg=COLORS["red"])

    def _export_excel(self):
        if not self.trades:
            messagebox.showwarning("提示","暂无交易记录可导出")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel文件","*.xlsx")],
            initialfile=f"交易复盘_{datetime.date.today()}.xlsx"
        )
        if not path: return
        try:
            export_to_excel(self.trades, path)
            messagebox.showinfo("导出成功", f"Excel已保存：\n{path}")
            if messagebox.askyesno("打开文件","是否立即打开Excel？"):
                import subprocess, sys
                if sys.platform == "darwin":
                    subprocess.run(["open", path])
                elif sys.platform == "win32":
                    os.startfile(path)
                else:
                    subprocess.run(["xdg-open", path])
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def _get_api_key(self):
        if API_KEY_FILE.exists():
            return API_KEY_FILE.read_text(encoding="utf-8").strip()
        return self.data.get("settings", {}).get("api_key", "")

    def _settings_dialog(self):
        dlg = tk.Toplevel(self.root)
        dlg.title("设置")
        dlg.geometry("520x200")
        dlg.configure(bg=COLORS["bg"])
        dlg.grab_set()

        tk.Label(dlg, text="Claude API Key", bg=COLORS["bg"],
                 fg=COLORS["text"], font=("Helvetica",10,"bold")).pack(anchor="w", padx=20, pady=(20,4))
        tk.Label(dlg, text="获取地址：https://console.anthropic.com/  （用于图片识别功能）",
                 bg=COLORS["bg"], fg=COLORS["muted"], font=("Helvetica",8)).pack(anchor="w", padx=20)

        var = tk.StringVar(value=self._get_api_key())
        entry = tk.Entry(dlg, textvariable=var, bg=COLORS["input_bg"],
                         fg=COLORS["text"], insertbackground=COLORS["text"],
                         relief="flat", font=("Helvetica",10), show="•",
                         highlightthickness=1, highlightbackground=COLORS["border"],
                         width=50)
        entry.pack(padx=20, pady=8, fill="x", ipady=6)

        def save_key():
            key = var.get().strip()
            API_KEY_FILE.write_text(key, encoding="utf-8")
            self.data.setdefault("settings",{})["api_key"] = key
            save_data(self.data)
            messagebox.showinfo("✓","API Key 已保存")
            dlg.destroy()

        tk.Button(dlg, text="保存", command=save_key,
                  **self._btn_style(accent=True)).pack(pady=6)


# ── Entry point ───────────────────────────────────────────────────────────────

if __name__ == "__main__":
    root = tk.Tk()
    app  = TradeJournalApp(root)
    root.mainloop()
